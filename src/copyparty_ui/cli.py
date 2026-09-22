#!/usr/bin/env python3
"""
copyparty_ui.cli — launcher for the CopyParty modern WebUI.

Usage:
    python -m copyparty_ui [--dir DIR] [options]

Examples:
    python -m copyparty_ui                          # serve %USERPROFILE%\\Downloads
    python -m copyparty_ui --dir C:\\files            # serve a specific directory
    python -m copyparty_ui --dir E:\\ --port 8080     # custom port on an external drive
"""

import os
import sys
import re
import json
import sqlite3
import argparse
import subprocess
import threading
import webbrowser
import http.server
import socket
import time
import urllib.parse
import urllib.request
import io
import base64
from contextlib import ExitStack
from importlib import resources
from pathlib import Path

PACKAGE = "copyparty_ui"

# Per-user, per-machine app data — never write runtime state into the
# installed package directory (site-packages may not be writable, and
# would collide across venvs/users).
APPDATA_DIR = Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "copyparty_ui"

# Served directory defaults to the current user's Downloads folder.
DEFAULT_SERVE_DIR = Path(os.environ.get("USERPROFILE", str(Path.home()))) / "Downloads"

COPYPARTY_PORT = 3923
UI_PORT = 3924

# favicon.ico embedded as base64 — eliminates the need for a separate file.
# Regenerate with: python3 -c "import base64; print(base64.b64encode(open('favicon.ico','rb').read()).decode())"
_FAVICON_B64 = (
    "AAABAAMAEBAAAAEAIACZAAAANgAAACAgAAABACAA4AAAAM8AAAAwMAAAAQAgAC0BAACvAQAAiVBORw0K"
    "GgoAAAANSUhEUgAAABAAAAAQCAYAAAAf8/9hAAAAYElEQVR4nGNkYGBgUE1+/Z+BDHB7rigjI7ma"
    "YYCJEs30MeDWHJEBdMEtqO34XMFEMxfcQrMVlysYCaUDkEa1lDc45RmRDSAU4sgAZigT0TqIcQED"
    "GWAQJGVQliRXM0gvAAhPHrINKHZpAAAAAElFTkSuQmCCiVBORw0KGgoAAAANSUhEUgAAACAAAAAgCAYA"
    "AABzenr0AAAAp0lEQVR4nO2X0Q2DMAxE7VNHgBHKYnQsuhgdge5QRKVUVgRCONQnNX1f/Dh3vhgZ"
    "VAzXfnpJAI97q+kZ0eK5FqLFcxPKELeAKf438BsJjEPDNVAKzui+JAUIGXgL8669KUDIwFO01a0n"
    "BQgZHC3Y6/JoClqyDa1Yd3u6zoCQ0eq/By7fWjJrrM0JhAyqN6DVvwWofgZgfxSjWbTfV8AwkTQ/"
    "MxBpwmrNF5Y0/D1wi9AAAAAASUVORK5CYIKJUE5HDQoaCgAAAA1JSERSAAAAMAAAADAIBgAAAFcC+Yc"
    "AAAD0SURBVHic7ZrRDYIwAESvF0eQEXQxHUsWsyPoDvojBkgLmrTSl/g++bhe7xqg0KAMh9PtoYaI"
    "fRdS10PrxtcmYpL5lEeTzKe8mmZ+YPA8WUJEAjH9MfgGLDgWHAuOBce1B7he9lX1LTj+Rfo1W7Dg"
    "uJbwPPVaLVhwXEM0l3aNFiw4Li24lnLpFiw4Lin2abolW7DguJTQt6mWaiHU2lKmDB7P9+LjWHAs"
    "OOH/VWJjLDi7FraFS6zdufANWHAsOBac8H8ObIwFx4JjwbHgOPcHnEDsu8BvQAvnEFomvjy/GyBN"
    "Io68TpYQYRJx5jFruLVXjFy4Txv1S0KZPdOEAAAAAElFTkSuQmCC"
)
_FAVICON_BYTES = base64.b64decode(_FAVICON_B64)

# Text file extensions to index for content search
TEXT_EXTS = {
    "txt", "md", "csv", "log", "json", "yaml", "yml", "xml", "html", "htm",
    "ini", "cfg", "conf", "toml", "sh", "bat", "ps1", "py", "js", "ts",
    "tsx", "jsx", "css", "scss", "less", "sql", "r", "rst", "tsv", "nfo",
}

# Skip paths matching this pattern during content indexing
SKIP_RE = re.compile(r"[\\/](node_modules|\.git|__pycache__|\.hist|\.hg|\.svn)[\\/]", re.IGNORECASE)

MAX_FILE_BYTES = 10 * 1024 * 1024  # skip files larger than 10 MB


# ---------------------------------------------------------------------------
# Content search index
# ---------------------------------------------------------------------------

class ContentIndex:
    def __init__(self, serve_dir: Path):
        self.serve_dir = serve_dir
        self.db_path = APPDATA_DIR / ".hist" / "content_fts.db"
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()
        self._conn: sqlite3.Connection | None = None
        self._ready = False

    def _connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(str(self.db_path), check_same_thread=False)
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("""
            CREATE VIRTUAL TABLE IF NOT EXISTS files USING fts5(
                path UNINDEXED,
                content,
                tokenize='unicode61'
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS meta (
                path TEXT PRIMARY KEY,
                mtime REAL,
                size INTEGER
            )
        """)
        conn.commit()
        return conn

    def _should_index(self, path: Path) -> bool:
        if SKIP_RE.search(str(path)):
            return False
        ext = path.suffix.lstrip(".").lower()
        return ext in TEXT_EXTS

    def build(self):
        """Full scan in background thread. Safe to call repeatedly."""
        conn = self._connect()
        self._conn = conn

        print("[content-index] Scanning for text files...")
        t0 = time.time()
        added = updated = skipped = 0

        try:
            for root, dirs, files in os.walk(self.serve_dir):
                # Prune dirs in-place so os.walk skips them
                dirs[:] = [d for d in dirs
                           if not SKIP_RE.search(os.path.join(root, d) + os.sep)]

                for fname in files:
                    fpath = Path(root) / fname
                    if not self._should_index(fpath):
                        continue

                    try:
                        st = fpath.stat()
                    except OSError:
                        continue

                    if st.st_size > MAX_FILE_BYTES:
                        skipped += 1
                        continue

                    rel = str(fpath.relative_to(self.serve_dir)).replace("\\", "/")

                    with self._lock:
                        row = conn.execute(
                            "SELECT mtime, size FROM meta WHERE path=?", (rel,)
                        ).fetchone()

                    if row and abs(row[0] - st.st_mtime) < 0.5 and row[1] == st.st_size:
                        continue  # unchanged

                    try:
                        text = fpath.read_text(encoding="utf-8", errors="ignore")
                    except OSError:
                        continue

                    with self._lock:
                        conn.execute("DELETE FROM files WHERE path=?", (rel,))
                        conn.execute("INSERT INTO files(path, content) VALUES(?,?)", (rel, text))
                        conn.execute(
                            "INSERT OR REPLACE INTO meta(path, mtime, size) VALUES(?,?,?)",
                            (rel, st.st_mtime, st.st_size),
                        )
                        if row:
                            updated += 1
                        else:
                            added += 1

            with self._lock:
                conn.commit()

        except Exception as ex:
            print(f"[content-index] Error during scan: {ex}")

        elapsed = time.time() - t0
        print(f"[content-index] Done in {elapsed:.1f}s - {added} added, {updated} updated, {skipped} skipped (too large)")
        self._ready = True

    def search(self, query: str, limit: int = 50):
        if not self._ready or self._conn is None:
            return {"hits": [], "indexing": True}

        # Sanitize: FTS5 MATCH syntax — escape quotes, wrap in ""
        safe = query.replace('"', '""')
        fts_query = f'"{safe}"'

        try:
            with self._lock:
                rows = self._conn.execute(
                    """
                    SELECT path,
                           snippet(files, 1, '>>>',  '<<<', '…', 20)
                    FROM files
                    WHERE files MATCH ?
                    ORDER BY rank
                    LIMIT ?
                    """,
                    (fts_query, limit),
                ).fetchall()
        except sqlite3.OperationalError as ex:
            return {"hits": [], "error": str(ex)}

        return {
            "hits": [{"path": r[0], "snippet": r[1]} for r in rows],
            "total": len(rows),
            "truncated": len(rows) == limit,
        }


# ---------------------------------------------------------------------------
# HTTP server (serves UI + content-search endpoint)
# ---------------------------------------------------------------------------

def render_pptx_slides(file_path: str, scale: float = 1.5) -> list:
    """
    Render each slide of a PPTX file to a base64 PNG using python-pptx + Pillow.
    Returns a list of data-URI strings ("data:image/png;base64,...").

    NOTE: This is an approximate renderer. It handles:
      - Solid background fills
      - Embedded images (PNG/JPG)
      - Text boxes positioned at correct coordinates
    Complex features (gradients, SmartArt, charts, theme fonts) are not rendered.
    """
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        from pptx.dml.color import RGBColor
        from PIL import Image, ImageDraw, ImageFont
    except ImportError as e:
        raise RuntimeError(f"Missing dependency: {e}. Install with: pip install copyparty-ui[pptx]")

    EMU = 914400.0

    prs = Presentation(file_path)
    W_emu = prs.slide_width
    H_emu = prs.slide_height
    W_px = int(W_emu / EMU * 96 * scale)
    H_px = int(H_emu / EMU * 96 * scale)

    def emu_to_px(emu):
        return int(emu / EMU * 96 * scale)

    def get_rgb(color_obj):
        try:
            if color_obj and color_obj.type is not None:
                rgb = color_obj.rgb
                return (rgb.red, rgb.green, rgb.blue)
        except Exception:
            pass
        return None

    # Try to load a basic font; fall back to default
    try:
        font_sm = ImageFont.truetype("arial.ttf", int(14 * scale))
        font_md = ImageFont.truetype("arial.ttf", int(18 * scale))
        font_lg = ImageFont.truetype("arial.ttf", int(24 * scale))
    except Exception:
        font_sm = font_md = font_lg = ImageFont.load_default()

    slides_b64 = []

    prs_part = prs.part
    slide_rels = [r for r in prs_part.rels.values()
                  if 'slide' in r.reltype.lower()
                  and 'layout' not in r.reltype.lower()
                  and 'master' not in r.reltype.lower()]

    for rel in slide_rels:
        slide_part = rel.target_part

        # Background colour
        bg_color = (255, 255, 255)
        try:
            bg_xml = slide_part._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
            if bg_xml is not None:
                srgb = bg_xml.find('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
                if srgb is not None:
                    val = srgb.get('val', 'ffffff')
                    r2 = int(val[0:2], 16); g2 = int(val[2:4], 16); b2 = int(val[4:6], 16)
                    bg_color = (r2, g2, b2)
        except Exception:
            pass

        img = Image.new('RGB', (W_px, H_px), bg_color)
        draw = ImageDraw.Draw(img)

        # Iterate shapes via XML
        NS = {
            'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
            'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
            'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
            'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture',
        }

        sp_tree = slide_part._element.find('.//p:spTree', NS)
        if sp_tree is None:
            sp_tree = slide_part._element

        for sp in sp_tree:
            tag = sp.tag.split('}')[-1] if '}' in sp.tag else sp.tag

            # Get transform (position + size)
            xfrm = sp.find('.//a:xfrm', NS)
            if xfrm is None:
                continue
            off = xfrm.find('a:off', NS)
            ext = xfrm.find('a:ext', NS)
            if off is None or ext is None:
                continue

            x = emu_to_px(int(off.get('x', 0)))
            y = emu_to_px(int(off.get('y', 0)))
            w = emu_to_px(int(ext.get('cx', 0)))
            h = emu_to_px(int(ext.get('cy', 0)))

            if w <= 0 or h <= 0:
                continue

            # Picture shapes
            if tag in ('pic', 'sp') or sp.find('.//a:blip', NS) is not None:
                blip = sp.find('.//a:blip', NS)
                if blip is not None:
                    rId = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    if rId and rId in slide_part.rels:
                        try:
                            img_part = slide_part.rels[rId].target_part
                            img_bytes = img_part.blob
                            sub = Image.open(io.BytesIO(img_bytes)).convert('RGBA')
                            sub = sub.resize((max(1, w), max(1, h)), Image.LANCZOS)
                            img.paste(sub, (x, y), sub)
                        except Exception:
                            pass
                    continue

            # Text shapes
            txBody = sp.find('.//p:txBody', NS) or sp.find('.//a:txBody', NS)
            if txBody is None:
                txBody = sp.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}txBody')
            if txBody is not None:
                lines = []
                for para in txBody.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}p'):
                    parts = []
                    for r_el in para.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}t'):
                        if r_el.text:
                            parts.append(r_el.text)
                    line = ''.join(parts)
                    if line.strip():
                        lines.append(line)

                if lines:
                    # Choose font size by box height
                    font = font_lg if h > emu_to_px(int(1.5 * EMU)) else font_md if h > emu_to_px(int(0.8 * EMU)) else font_sm
                    text_color = (30, 30, 30)

                    # Solid fill for text box background
                    fill = sp.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill/{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
                    if fill is not None:
                        val = fill.get('val', '')
                        if val:
                            try:
                                br = int(val[0:2], 16); bg = int(val[2:4], 16); bb = int(val[4:6], 16)
                                draw.rectangle([x, y, x + w, y + h], fill=(br, bg, bb))
                            except Exception:
                                pass

                    ty = y + int(h * 0.1)
                    for line in lines:
                        if ty > y + h:
                            break
                        draw.text((x + 6, ty), line, fill=text_color, font=font)
                        ty += int(font.size * 1.3) if hasattr(font, 'size') else 18

        # Encode to base64 PNG
        buf = io.BytesIO()
        img.save(buf, format='PNG', optimize=True)
        b64 = base64.b64encode(buf.getvalue()).decode()
        slides_b64.append(f'data:image/png;base64,{b64}')

    return slides_b64


class SharesDB:
    """Read-only access to copyparty's shares SQLite database."""
    def __init__(self):
        self._path = Path(os.path.expandvars(r"%APPDATA%")) / "copyparty" / "shares.db"
        self._lock = threading.Lock()

    # copyparty stores perms as a compact string e.g. "rg." where each char maps to a permission
    _PERM_MAP = {'r': 'read', 'w': 'write', 'm': 'move', 'd': 'delete', 'g': 'get'}

    def list(self) -> list:
        if not self._path.exists():
            return []
        try:
            with self._lock:
                conn = sqlite3.connect(str(self._path), check_same_thread=False)
                rows = conn.execute("SELECT k, pw, vp, pr, st, un, t0, t1 FROM sh ORDER BY t0 DESC").fetchall()
                result = []
                for r in rows:
                    skey, pw, vp, pr, _st, un, t0, t1 = r
                    fns = [x[0] for x in conn.execute("SELECT vp FROM sf WHERE k=?", (skey,)).fetchall()]
                    perms = [self._PERM_MAP[c] for c in pr if c in self._PERM_MAP]
                    result.append({
                        "key": skey,
                        "vp": vp,
                        "files": fns,
                        "perms": perms,
                        "created": t0,
                        "expires": t1,
                        "protected": bool(pw),
                    })
                conn.close()
                return result
        except Exception as ex:
            print(f"[shares-db] {ex}")
            return []


_shares_db = SharesDB()


# ---------------------------------------------------------------------------
# Labels database
# ---------------------------------------------------------------------------

LABEL_COLORS = [
    "#ef4444", "#f97316", "#eab308", "#22c55e", "#14b8a6",
    "#3b82f6", "#8b5cf6", "#ec4899", "#6b7280", "#0ea5e9",
]

class LabelsDB:
    """User-defined file labels stored in %LOCALAPPDATA%\\copyparty_ui\\.hist\\labels.db."""

    def __init__(self):
        self._path = APPDATA_DIR / ".hist" / "labels.db"
        self._path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()
        self._conn = self._connect()

    def _connect(self):
        conn = sqlite3.connect(str(self._path), check_same_thread=False)
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("""
            CREATE TABLE IF NOT EXISTS labels (
                name TEXT PRIMARY KEY,
                color TEXT NOT NULL
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS file_labels (
                path TEXT NOT NULL,
                label TEXT NOT NULL,
                PRIMARY KEY (path, label),
                FOREIGN KEY (label) REFERENCES labels(name) ON DELETE CASCADE
            )
        """)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_fl_label ON file_labels(label)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_fl_path  ON file_labels(path)")
        conn.commit()
        return conn

    def get_labels(self) -> list:
        with self._lock:
            rows = self._conn.execute("SELECT name, color FROM labels ORDER BY name").fetchall()
            # Include file count per label
            counts = dict(self._conn.execute(
                "SELECT label, COUNT(*) FROM file_labels GROUP BY label").fetchall())
            return [{"name": r[0], "color": r[1], "count": counts.get(r[0], 0)} for r in rows]

    def create_label(self, name: str, color: str) -> dict:
        name = name.strip()
        if not name:
            raise ValueError("Label name cannot be empty")
        if len(name) > 50:
            raise ValueError("Label name too long (max 50 chars)")
        with self._lock:
            self._conn.execute(
                "INSERT OR IGNORE INTO labels(name, color) VALUES(?,?)", (name, color))
            self._conn.commit()
        return {"name": name, "color": color, "count": 0}

    def update_label(self, name: str, new_name: str | None, color: str | None) -> None:
        with self._lock:
            if new_name and new_name != name:
                self._conn.execute(
                    "UPDATE file_labels SET label=? WHERE label=?", (new_name, name))
                self._conn.execute(
                    "UPDATE labels SET name=? WHERE name=?", (new_name, name))
            if color:
                target = new_name or name
                self._conn.execute(
                    "UPDATE labels SET color=? WHERE name=?", (color, target))
            self._conn.commit()

    def delete_label(self, name: str) -> None:
        with self._lock:
            self._conn.execute("DELETE FROM file_labels WHERE label=?", (name,))
            self._conn.execute("DELETE FROM labels WHERE name=?", (name,))
            self._conn.commit()

    def get_file_labels(self, path: str) -> list:
        with self._lock:
            rows = self._conn.execute(
                "SELECT l.name, l.color FROM labels l "
                "JOIN file_labels fl ON fl.label=l.name "
                "WHERE fl.path=? ORDER BY l.name", (path,)).fetchall()
            return [{"name": r[0], "color": r[1]} for r in rows]

    def get_bulk_labels(self, paths: list) -> dict:
        """Return {path: [{name, color}]} for multiple paths at once."""
        if not paths:
            return {}
        with self._lock:
            placeholders = ",".join("?" * len(paths))
            rows = self._conn.execute(
                f"SELECT fl.path, l.name, l.color FROM file_labels fl "
                f"JOIN labels l ON l.name=fl.label "
                f"WHERE fl.path IN ({placeholders}) ORDER BY fl.path, l.name",
                paths).fetchall()
        result: dict = {p: [] for p in paths}
        for path, lname, lcolor in rows:
            result[path].append({"name": lname, "color": lcolor})
        return result

    def set_file_label(self, path: str, label: str, assigned: bool) -> None:
        with self._lock:
            # Ensure label exists
            exists = self._conn.execute(
                "SELECT 1 FROM labels WHERE name=?", (label,)).fetchone()
            if not exists:
                raise ValueError(f"Label {label!r} does not exist")
            if assigned:
                self._conn.execute(
                    "INSERT OR IGNORE INTO file_labels(path, label) VALUES(?,?)",
                    (path, label))
            else:
                self._conn.execute(
                    "DELETE FROM file_labels WHERE path=? AND label=?",
                    (path, label))
            self._conn.commit()

    def set_bulk_labels(self, paths: list, label: str, assigned: bool) -> None:
        with self._lock:
            exists = self._conn.execute(
                "SELECT 1 FROM labels WHERE name=?", (label,)).fetchone()
            if not exists:
                raise ValueError(f"Label {label!r} does not exist")
            for path in paths:
                if assigned:
                    self._conn.execute(
                        "INSERT OR IGNORE INTO file_labels(path, label) VALUES(?,?)",
                        (path, label))
                else:
                    self._conn.execute(
                        "DELETE FROM file_labels WHERE path=? AND label=?",
                        (path, label))
            self._conn.commit()

    def get_files_with_label(self, label: str) -> list:
        with self._lock:
            rows = self._conn.execute(
                "SELECT path FROM file_labels WHERE label=? ORDER BY path",
                (label,)).fetchall()
            return [r[0] for r in rows]


_labels_db = LabelsDB()


def make_handler(html_bytes: bytes, index: ContentIndex, copyparty_port: int):
    class Handler(http.server.BaseHTTPRequestHandler):
        def _json(self, data, status=200):
            body = json.dumps(data).encode()
            self.send_response(status)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(body)

        def do_GET(self):
            parsed = urllib.parse.urlparse(self.path)
            params = urllib.parse.parse_qs(parsed.query)

            if parsed.path == "/content-search":
                q = (params.get("q") or [""])[0].strip()
                limit = int((params.get("n") or ["50"])[0])
                body_data = {"hits": [], "total": 0} if not q else index.search(q, limit)
                self._json(body_data)
                return

            if parsed.path == "/favicon.ico":
                self.send_response(200)
                self.send_header("Content-Type", "image/x-icon")
                self.send_header("Content-Length", str(len(_FAVICON_BYTES)))
                self.send_header("Cache-Control", "public, max-age=86400")
                self.end_headers()
                self.wfile.write(_FAVICON_BYTES)
                return

            if parsed.path == "/api/shares":
                self._json(_shares_db.list())
                return

            if parsed.path == "/api/labels":
                self._json(_labels_db.get_labels())
                return

            if parsed.path == "/api/file-labels":
                path = (params.get("path") or [""])[0].strip()
                if not path:
                    self._json({"error": "missing path"}, 400)
                    return
                self._json(_labels_db.get_file_labels(path))
                return

            if parsed.path == "/api/bulk-labels":
                raw = (params.get("paths") or [""])[0]
                paths = [p for p in raw.split(",") if p]
                self._json(_labels_db.get_bulk_labels(paths))
                return

            if parsed.path == "/api/scan":
                try:
                    r = urllib.request.urlopen(f"http://127.0.0.1:{copyparty_port}/?scan", timeout=10)
                    self._json({"ok": True, "status": r.status})
                except Exception as ex:
                    self._json({"ok": False, "error": str(ex)}, 500)
                return

            if parsed.path == "/api/pptx-slides":
                file_url = (params.get("url") or [""])[0].strip()
                if not file_url:
                    self._json({"error": "missing url param"}, 400)
                    return
                try:
                    # Download the file from copyparty
                    resp = urllib.request.urlopen(file_url, timeout=30)
                    pptx_bytes = resp.read()
                    # Write to a temp file (python-pptx needs a path)
                    import tempfile
                    with tempfile.NamedTemporaryFile(suffix='.pptx', delete=False) as tf:
                        tf.write(pptx_bytes)
                        tmp_path = tf.name
                    try:
                        slides = render_pptx_slides(tmp_path)
                        self._json({"slides": slides})
                    finally:
                        try: os.unlink(tmp_path)
                        except: pass
                except Exception as ex:
                    self._json({"error": str(ex)}, 500)
                return

            if parsed.path == "/dl":
                # Proxy a single file from copyparty with Content-Disposition: attachment.
                # ?path=/some/file.py  (virtual path as copyparty knows it)
                file_path = (params.get("path") or [""])[0]
                if not file_path:
                    self._json({"error": "missing path"}, 400)
                    return
                target = f"http://127.0.0.1:{copyparty_port}{file_path}"
                filename = file_path.rstrip("/").rsplit("/", 1)[-1]
                try:
                    upstream = urllib.request.urlopen(target, timeout=30)
                    ctype = upstream.headers.get("Content-Type", "application/octet-stream")
                    clength = upstream.headers.get("Content-Length")
                    self.send_response(200)
                    self.send_header("Content-Type", ctype)
                    safe_name = filename.encode("ascii", "replace").decode()
                    self.send_header(
                        "Content-Disposition",
                        f'attachment; filename="{safe_name}"; filename*=UTF-8\'\'{urllib.parse.quote(filename)}'
                    )
                    if clength:
                        self.send_header("Content-Length", clength)
                    self.end_headers()
                    while True:
                        chunk = upstream.read(65536)
                        if not chunk:
                            break
                        self.wfile.write(chunk)
                except Exception as ex:
                    self._json({"error": str(ex)}, 502)
                return

            if parsed.path == "/zip":
                # Proxy a selective ZIP from copyparty.
                # ?dir=/some/path/&name=file1&name=file2&...
                # POSTs to copyparty with act=zip + files=<newline-separated names>.
                dir_path = (params.get("dir") or [""])[0]
                names = params.get("name") or []
                if not dir_path or not names:
                    self._json({"error": "missing dir or name params"}, 400)
                    return
                target = f"http://127.0.0.1:{copyparty_port}{dir_path}?zip=crc"
                body = ("--b\r\nContent-Disposition: form-data; name=\"act\"\r\n\r\nzip\r\n"
                        "--b\r\nContent-Disposition: form-data; name=\"files\"\r\n\r\n"
                        + "\n".join(names)
                        + "\r\n--b--\r\n")
                body_bytes = body.encode()
                req = urllib.request.Request(
                    target,
                    data=body_bytes,
                    method="POST",
                    headers={"Content-Type": "multipart/form-data; boundary=b",
                             "Content-Length": str(len(body_bytes))},
                )
                zip_name = dir_path.strip("/").split("/")[-1] or "download"
                try:
                    upstream = urllib.request.urlopen(req, timeout=60)
                    self.send_response(200)
                    self.send_header("Content-Type", "application/zip")
                    self.send_header(
                        "Content-Disposition",
                        f'attachment; filename="{zip_name}.zip"'
                    )
                    self.end_headers()
                    while True:
                        chunk = upstream.read(65536)
                        if not chunk:
                            break
                        self.wfile.write(chunk)
                except Exception as ex:
                    self._json({"error": str(ex)}, 502)
                return

            # All other paths → serve index.html
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(html_bytes)))
            self.end_headers()
            self.wfile.write(html_bytes)

        def do_POST(self):
            parsed = urllib.parse.urlparse(self.path)
            length = int(self.headers.get("Content-Length", 0))
            body = json.loads(self.rfile.read(length)) if length else {}

            if parsed.path == "/api/labels":
                # Create label: {name, color}
                try:
                    label = _labels_db.create_label(
                        body.get("name", ""),
                        body.get("color", LABEL_COLORS[0]))
                    self._json(label, 201)
                except ValueError as ex:
                    self._json({"error": str(ex)}, 400)
                return

            if parsed.path == "/api/labels/update":
                # Update label: {name, new_name?, color?}
                try:
                    _labels_db.update_label(
                        body.get("name", ""),
                        body.get("new_name"),
                        body.get("color"))
                    self._json({"ok": True})
                except Exception as ex:
                    self._json({"error": str(ex)}, 400)
                return

            if parsed.path == "/api/labels/delete":
                # Delete label: {name}
                try:
                    _labels_db.delete_label(body.get("name", ""))
                    self._json({"ok": True})
                except Exception as ex:
                    self._json({"error": str(ex)}, 400)
                return

            if parsed.path == "/api/file-labels":
                # Assign/unassign: {path, label, assigned}
                try:
                    _labels_db.set_file_label(
                        body["path"], body["label"], bool(body.get("assigned", True)))
                    self._json({"ok": True})
                except Exception as ex:
                    self._json({"error": str(ex)}, 400)
                return

            if parsed.path == "/api/bulk-labels":
                # Bulk assign: {paths: [...], label, assigned}
                try:
                    _labels_db.set_bulk_labels(
                        body.get("paths", []),
                        body["label"],
                        bool(body.get("assigned", True)))
                    self._json({"ok": True})
                except Exception as ex:
                    self._json({"error": str(ex)}, 400)
                return

            self._json({"error": "not found"}, 404)

        def do_OPTIONS(self):
            self.send_response(204)
            self.send_header("Access-Control-Allow-Origin", "*")
            self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
            self.send_header("Access-Control-Allow-Headers", "Content-Type")
            self.end_headers()

        def log_message(self, *args):
            pass

    return Handler


def serve_ui(port: int, copyparty_port: int, index: ContentIndex, ui_html_path: Path):
    html = ui_html_path.read_text(encoding="utf-8")
    inject = (
        f'<script>'
        f'window.__COPYPARTY_BASE__="http://"+window.location.hostname+":{copyparty_port}";'
        f'window.__UI_BASE__="http://"+window.location.hostname+":{port}";'
        f'</script>'
    )
    html = html.replace("<head>", f"<head>{inject}", 1)
    html_bytes = html.encode("utf-8")

    Handler = make_handler(html_bytes, index, copyparty_port)
    server = http.server.HTTPServer(("0.0.0.0", port), Handler)
    server.serve_forever()


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def find_python():
    return sys.executable


def main():
    parser = argparse.ArgumentParser(
        description="CopyParty modern UI - a Google Drive-style frontend for copyparty.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=f"""
examples:
  python -m copyparty_ui                          serve {DEFAULT_SERVE_DIR}
  python -m copyparty_ui --dir X:\\                 serve a specific drive/folder
  python -m copyparty_ui --dir X:\\ --no-index      skip file hashing (recommended for large drives)
  python -m copyparty_ui --dir X:\\ --port 8080     custom copyparty backend port
  python -m copyparty_ui --dir X:\\ --ui-port 8081  custom UI port
  python -m copyparty_ui --dir X:\\ --no-browser    don't open browser on start

notes:
  --no-index skips the startup file-hash scan (-e2ds -> -e2d) and the
  built-in full-text content search. Move and rename still work.
  Strongly recommended for large drives (100GB+) - without it copyparty
  hashes every file on startup which can take hours and block all requests.
        """,
    )
    parser.add_argument("--dir", dest="directory", default=None,
                        help=f"Directory to serve (default: {DEFAULT_SERVE_DIR})")
    parser.add_argument("-p", "--port", type=int, default=COPYPARTY_PORT,
                        help=f"copyparty backend port (default: {COPYPARTY_PORT})")
    parser.add_argument("--ui-port", type=int, default=UI_PORT,
                        help=f"UI server port (default: {UI_PORT})")
    parser.add_argument("--no-browser", action="store_true",
                        help="Don't open browser automatically on start")
    parser.add_argument("--no-index", action="store_true",
                        help="Skip the startup file-hash scan and full-text content indexing. "
                             "Move/rename still work. Recommended for large drives - the default "
                             "scan hashes every file on startup which can take hours and block requests.")
    args = parser.parse_args()

    serve_dir = Path(args.directory).resolve() if args.directory else DEFAULT_SERVE_DIR.resolve()
    if not serve_dir.exists():
        print(f"ERROR: Directory not found: {serve_dir}")
        sys.exit(1)

    with ExitStack() as stack:
        pkg = resources.files(PACKAGE)
        sfx_res = pkg / "vendor" / "copyparty-sfx.py"
        ui_res = pkg / "webui" / "index.html"
        sfx_path = stack.enter_context(resources.as_file(sfx_res))
        ui_path = stack.enter_context(resources.as_file(ui_res))

        if not sfx_path.exists():
            print(f"ERROR: copyparty-sfx.py not found at {sfx_path}")
            sys.exit(1)

        if not ui_path.exists():
            print(f"ERROR: index.html not found at {ui_path}")
            sys.exit(1)

        local_ip = socket.gethostbyname(socket.gethostname())
        print(f"[copyparty-ui] Serving:  {serve_dir}")
        print(f"[copyparty-ui] UI:       http://localhost:{args.ui_port}  |  http://{local_ip}:{args.ui_port}")
        print(f"[copyparty-ui] Backend:  http://localhost:{args.port}  |  http://{local_ip}:{args.port}")

        # Build content index in background (skipped with --no-index)
        index = ContentIndex(serve_dir)
        if args.no_index:
            print("[copyparty-ui] Content indexing disabled (--no-index)")
        else:
            threading.Thread(target=index.build, daemon=True).start()

        # Start UI server
        ui_thread = threading.Thread(
            target=serve_ui, args=(args.ui_port, args.port, index, ui_path), daemon=True)
        ui_thread.start()

        # Open browser
        if not args.no_browser:
            def _open():
                time.sleep(1.5)
                webbrowser.open(f"http://127.0.0.1:{args.ui_port}")
            threading.Thread(target=_open, daemon=True).start()

        # Start copyparty
        cmd = [
            find_python(),
            str(sfx_path),
            "-p", str(args.port),
            "-v", f"{serve_dir}::rwmdA",
            "--no-reload",
            "--acao", "*",
            "--acam", "GET,HEAD,PUT,DELETE,POST,OPTIONS",
            "--no-idx", r"[\\/](node_modules|\.git|__pycache__|\.hist)[\\/]",
            "--shr", "shares",
            "--shr-who", "a",
        ]
        if args.no_index:
            cmd.append("-e2d")   # broker on (required for move/rename), no startup scan
        else:
            cmd.append("-e2ds")  # broker on + scan all writable folders on startup

        print("[copyparty-ui] Starting copyparty...")
        try:
            subprocess.run(cmd)
        except KeyboardInterrupt:
            print("\n[copyparty-ui] Stopped.")


if __name__ == "__main__":
    main()
